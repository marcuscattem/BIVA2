#!/usr/bin/env python3
"""Build the BIVA2 reference database from the supplied literature workbook."""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = Path.home() / "Downloads" / "REF_BIVA_GB_MARCUS 2.xlsx"


def clean(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def ascii_key(value) -> str:
    text = unicodedata.normalize("NFD", clean(value))
    return "".join(ch for ch in text if unicodedata.category(ch) != "Mn").lower()


def doi_from(text) -> str:
    match = re.search(r"10\.\d{4,9}/[-._;()/:A-Z0-9]+", clean(text), re.I)
    return match.group(0).rstrip(".,;") if match else ""


def infer_sex(legend, raw="") -> str:
    value = clean(raw).lower()
    if value in {"m", "male"}:
        return "M"
    if value in {"f", "female"}:
        return "F"
    if re.search(r"males?\s*\+\s*females?|both sexes|mixed", clean(legend), re.I):
        return "B"
    if re.search(r"\bfemales?\b|\bwomen\b", clean(legend), re.I):
        return "F"
    if re.search(r"\bmales?\b|\bmen\b", clean(legend), re.I):
        return "M"
    return "—"


COUNTRIES = (
    (r"\bUSA\b|United States|American|NHANES", "USA"),
    (r"Brazil|Brazilian", "Brazil"),
    (r"Ital(?:y|ian)|Sardinia|Sardinian", "Italy"),
    (r"Mexic", "Mexico"),
    (r"Spain|Spanish", "Spain"),
    (r"Portugal|Portuguese", "Portugal"),
    (r"German", "Germany"),
    (r"Japan|Japanese", "Japan"),
    (r"Korea|Korean", "South Korea"),
    (r"Poland|Polish", "Poland"),
    (r"Slovak", "Slovakia"),
    (r"Cuba|Cuban", "Cuba"),
    (r"India|Indian", "India"),
    (r"Taiwan", "Taiwan"),
    (r"Leban", "Lebanon"),
    (r"Netherlands|Dutch", "Netherlands"),
    (r"France|French", "France"),
    (r"Turkey|Turkish", "Turkey"),
)


def infer_country(legend) -> str:
    for pattern, country in COUNTRIES:
        if re.search(pattern, clean(legend), re.I):
            return country
    return "—"


def infer_age(legend) -> str:
    text = clean(legend)
    patterns = (
        (r"(\d+(?:\.\d+)?)\s*(?:-|–|to|age)\s*(\d+(?:\.\d+)?)\s*(?:y|yr|years?)\b", lambda m: f"{m[1]}–{m[2]} y"),
        (r"([<>≥≤]\s*\d+(?:\.\d+)?)\s*(?:y|yr|years?)\b", lambda m: f"{m[1]} y"),
        (r"mean\s*(\d+(?:\.\d+)?)\s*y", lambda m: f"mean {m[1]} y"),
        (r"(\d+(?:\.\d+)?)\s*±\s*(\d+(?:\.\d+)?)\s*y", lambda m: f"{m[1]} ± {m[2]} y"),
    )
    for pattern, formatter in patterns:
        match = re.search(pattern, text, re.I)
        if match:
            return formatter(match)
    for word, label in (("newborn", "Newborns"), ("neonate", "Newborns"), ("infant", "Infants"), ("children", "Children"), ("adolescent", "Adolescents"), ("older", "Older adults"), ("elderly", "Older adults"), ("adult", "Adults")):
        if word in text.lower():
            return label
    return "—"


def infer_population(legend, section="") -> str:
    text = ascii_key(f"{section} {legend}")
    if re.search(r"athlet|sport|soccer|football|rugby|volley|cycl|marathon|handball|basket|swim|triathlon", text):
        return "Sport"
    if re.search(r"doent|clinic|patient|disease|renal|dialy|obes|cancer|sarcopen|parkinson|dementia|hospital|esrd", text):
        return "Clinical"
    if re.search(r"saudav|healthy|reference|adult|child|adolesc|newborn|infant|elderly|older", text):
        return "Healthy"
    return "Other"


def is_article(text) -> bool:
    value = clean(text)
    return bool(value and re.search(r"\b(?:19|20)\d{2}\b", value) and not re.match(r"^(Males?|Females?|Elderly)", value, re.I))


def valid(n, sx, sy, correlation) -> bool:
    return n > 2 and sx > 0 and sy > 0 and abs(correlation) < 1


def source_from(value) -> str:
    text = ascii_key(value)
    if "ambas" in text:
        return "both"
    if "confianca" in text:
        return "confidence"
    if "tolerancia" in text:
        return "tolerance"
    return "none"


def base_record(*, metric, source, n, mx, sx, my, sy, correlation, legend, equipment, reference, note="", found_in="", ellipse_use="", source_sheet="", source_row=0, sex_raw="", extra=None):
    result = {
        "n": n,
        "r_h_mean": mx,
        "r_h_sd": sx,
        "xc_h_mean": my,
        "xc_h_sd": sy,
        "correlation": correlation,
        "sex": infer_sex(legend, sex_raw),
        "legend": clean(legend),
        "country": infer_country(legend),
        "age": infer_age(legend),
        "equipment": clean(equipment) or "—",
        "population": infer_population(legend),
        "reference": clean(reference),
        "doi": doi_from(reference),
        "source": source,
        "metric": metric,
        "note": clean(note),
        "found_in": clean(found_in),
        "ellipse_use": clean(ellipse_use),
        "source_sheet": source_sheet,
        "source_row": source_row,
        "provenance": [f"{source_sheet}!{source_row}"],
    }
    if extra:
        result.update(extra)
    return result


def parse_main(workbook, excluded):
    sheet = workbook["Referencias BIVA"]
    records, reference, section = [], "", "Other"
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        values = list(row) + [None] * 13
        legend = clean(values[8])
        core = all(isinstance(values[index], (int, float)) for index in (1, 2, 3, 4, 5, 6))
        if legend and not core:
            if legend.upper() == legend and not re.search(r"\d{4}", legend):
                section = legend
            elif is_article(legend):
                reference = legend
        if not core:
            continue
        n, mx, sx, my, sy, correlation = values[1:7]
        source = source_from(values[12])
        reason = ""
        if not valid(n, sx, sy, correlation):
            reason = "invalid N, SD, or correlation"
        elif source == "none":
            reason = "article did not provide a usable ellipse type"
        if reason:
            excluded.append({"sheet": sheet.title, "row": row_number, "legend": legend, "reason": reason})
            continue
        record = base_record(
            metric="classic", source=source, n=n, mx=mx, sx=sx, my=my, sy=sy, correlation=correlation,
            legend=legend, equipment=values[9], reference=reference, note=values[10], found_in=values[11],
            ellipse_use=values[12], source_sheet=sheet.title, source_row=row_number,
        )
        record["population"] = infer_population(legend, section)
        records.append(record)
    return records


def parse_specific(workbook, excluded):
    sheet = workbook["BIVA especifica (ohm.cm)"]
    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        values = list(row) + [None] * 10
        if not all(isinstance(values[index], (int, float)) for index in range(6)):
            continue
        n, mx, sx, my, sy, correlation = values[:6]
        legend = clean(values[6])
        if not valid(n, sx, sy, correlation):
            excluded.append({"sheet": sheet.title, "row": row_number, "legend": legend, "reason": "invalid N, SD, or correlation"})
            continue
        reference = re.split(r"\s+[–—-]\s+", legend, maxsplit=1)[0] if re.search(r"\b(?:19|20)\d{2}\b", legend) else ""
        records.append(base_record(
            metric="specific", source="both", n=n, mx=mx, sx=sx, my=my, sy=sy, correlation=correlation,
            legend=legend, equipment=values[7], reference=reference, note=values[8], found_in=values[9],
            ellipse_use="Confidence and tolerance derived from reported parameters", source_sheet=sheet.title, source_row=row_number,
        ))
    return records


def parse_piccoli(workbook, sheet_name, source, excluded):
    sheet = workbook[sheet_name]
    records, reference = [], ""
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        values = list(row) + [None] * 14
        for candidate in (values[0], values[8]):
            if is_article(candidate):
                reference = clean(candidate)
        if not all(isinstance(values[index], (int, float)) for index in range(7)):
            continue
        _, n, mx, sx, my, sy, correlation = values[:7]
        legend = clean(values[8])
        if not valid(n, sx, sy, correlation):
            excluded.append({"sheet": sheet.title, "row": row_number, "legend": legend, "reason": "invalid N, SD, or correlation"})
            continue
        records.append(base_record(
            metric="classic", source=source, n=n, mx=mx, sx=sx, my=my, sy=sy, correlation=correlation,
            legend=legend, equipment=values[9], reference=reference,
            ellipse_use="Tolerance" if source == "tolerance" else "Confidence",
            source_sheet=sheet.title, source_row=row_number, sex_raw=values[7],
        ))
    return records


def parse_h2(workbook, excluded):
    sheet = workbook["Bosy-Westphal 2005 (H2-R)"]
    reference = clean(sheet.cell(1, 1).value)
    sheet_note = clean(sheet.cell(2, 1).value)
    records = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
        values = list(row) + [None] * 13
        required = (1, 8, 9, 10, 11, 12)
        if not all(isinstance(values[index], (int, float)) for index in required):
            continue
        legend, n = clean(values[0]), values[1]
        mx, sx, my, sy, correlation = values[8], values[9], values[10], values[11], values[12]
        if not valid(n, sx, sy, correlation):
            excluded.append({"sheet": sheet.title, "row": row_number, "legend": legend, "reason": "invalid N, SD, or correlation"})
            continue
        records.append(base_record(
            metric="h2", source="confidence", n=n, mx=mx, sx=sx, my=my, sy=sy, correlation=correlation,
            legend=legend, equipment="Data Input BIA 2000-S (supine, 50 kHz)", reference=reference,
            note=sheet_note, found_in="Workbook table", ellipse_use="95% confidence",
            source_sheet=sheet.title, source_row=row_number,
            extra={
                "height_mean_m": values[2], "height_sd_m": values[3],
                "r_ohm_mean": values[4], "r_ohm_sd": values[5],
                "xc_ohm_mean": values[6], "xc_ohm_sd": values[7],
            },
        ))
    return records


def signature(record):
    return (
        record["metric"], record["sex"], int(record["n"]),
        round(float(record["r_h_mean"]), 6), round(float(record["r_h_sd"]), 6),
        round(float(record["xc_h_mean"]), 6), round(float(record["xc_h_sd"]), 6),
        round(float(record["correlation"]), 6),
    )


def richer(left, right, field):
    a, b = clean(left.get(field)), clean(right.get(field))
    return b if len(b) > len(a) else a


def merge_records(records):
    merged = {}
    for record in records:
        key = signature(record)
        if key not in merged:
            merged[key] = record
            continue
        current = merged[key]
        sources = {current["source"], record["source"]}
        current["source"] = "both" if "both" in sources or sources == {"confidence", "tolerance"} else current["source"]
        for field in ("legend", "reference", "equipment", "note", "found_in", "ellipse_use"):
            current[field] = richer(current, record, field)
        current["doi"] = current.get("doi") or record.get("doi") or doi_from(current["reference"])
        current["country"] = current.get("country") if current.get("country") != "—" else record.get("country", "—")
        current["age"] = current.get("age") if current.get("age") != "—" else record.get("age", "—")
        current["population"] = current.get("population") if current.get("population") != "Other" else record.get("population", "Other")
        current["provenance"] = list(dict.fromkeys(current.get("provenance", []) + record.get("provenance", [])))
    result = list(merged.values())
    result.sort(key=lambda item: (item["metric"], item["source"], item["population"], item["country"], item["sex"], item["legend"]))
    for index, record in enumerate(result, 1):
        record["code"] = f"ref-{index:04d}"
    return result


def write_outputs(records, paired):
    json_path = ROOT / "reference_populations_final.json"
    data_path = ROOT / "biva2-data.js"
    json_text = json.dumps(records, ensure_ascii=False, indent=2) + "\n"
    json_path.write_text(json_text, encoding="utf-8")
    js = "window.BIVA2_POPULATIONS = " + json_text.rstrip() + ";\n\n"
    js += "window.BIVA2_PAIRED_REFERENCES = " + json.dumps(paired, ensure_ascii=False, indent=2) + ";\n"
    data_path.write_text(js, encoding="utf-8")


def main():
    input_path = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else DEFAULT_INPUT
    workbook = load_workbook(input_path, data_only=True, read_only=True)
    excluded = []
    raw = []
    raw.extend(parse_main(workbook, excluded))
    raw.extend(parse_specific(workbook, excluded))
    raw.extend(parse_piccoli(workbook, "TOLERANCE PICOLLI", "tolerance", excluded))
    raw.extend(parse_piccoli(workbook, "CONFIDENCE PICOLLI", "confidence", excluded))
    raw.extend(parse_h2(workbook, excluded))
    records = merge_records(raw)
    paired = json.loads((ROOT / "paired_reference_populations.json").read_text(encoding="utf-8"))
    write_outputs(records, paired)
    summary = {
        "input": str(input_path),
        "raw_eligible": len(raw),
        "duplicates_removed": len(raw) - len(records),
        "final_records": len(records),
        "by_metric": Counter(item["metric"] for item in records),
        "by_source": Counter(item["source"] for item in records),
        "excluded_rows": len(excluded),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=dict))


if __name__ == "__main__":
    main()
